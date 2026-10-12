"""Extract plain text from uploaded CVs (PDF, DOCX, TXT)."""
from pathlib import Path


def read_cv(path, max_chars: int = 2500) -> str:
    """Return CV text (truncated to max_chars), or "" if missing/unreadable."""
    if not path:
        return ""
    p = Path(path)
    if not p.exists():
        return ""

    try:
        suffix = p.suffix.lower()
        if suffix == ".docx":
            import docx
            d = docx.Document(str(p))
            parts = [para.text for para in d.paragraphs if para.text.strip()]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
            text = "\n".join(parts)
        elif suffix == ".pdf":
            from pypdf import PdfReader
            text = "\n".join((page.extract_text() or "") for page in PdfReader(str(p)).pages)
        else:
            text = p.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return ""

    return text.strip()[:max_chars]
